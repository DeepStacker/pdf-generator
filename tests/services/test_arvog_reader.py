"""Characterization tests for the wide-sheet reader.

These pin what `convert_raw_to_tall` does with the original master layout —
nine loan columns, then one seven-column block per ornament, then seven more
loan columns. That layout is the one in production use, so these tests exist to
prove a refactor of the column mapping changes nothing about it.

They were written against the pre-refactor implementation and must keep passing
verbatim. If one of them has to change, the refactor broke something real.
"""

import openpyxl
import pandas as pd
import pytest

from audit_engine.services.arvog import _MAX_JEWELLERY_COLUMNS, ArvogService

LEADING = ["SR No.", "Branch", "Region", "Date", "Customer Name", "Loan No.",
           "Loan type", "Partner Name", "Customer ID"]
TRAILING = ["Loan Amount", "Auditor Name", "Appraiser Name", "Packets Location",
            "Cluster Manager", "Packets Number", "After audit packet number"]

# What one ornament occupies in the original layout, after its Jewellery{i} name.
ORNAMENT_ATTRS = ["Gross Wt.", "Stone Wt.", "Net Wt.", "Karat", "Purity % %",
                  "Net weight after Purity %"]


def _original_layout_headers(ornaments: int) -> list[str]:
    """The wide master exactly as branches send it."""
    headers = list(LEADING)
    for i in range(1, ornaments + 1):
        headers.append(f"Jewellery{i}")
        attrs = list(ORNAMENT_ATTRS)
        if i == _MAX_JEWELLERY_COLUMNS:
            # The 20th block ends with a differently-spelled purity column.
            attrs[-1] = "Final net weight after purity %"
        headers.extend(attrs)
    headers.extend(TRAILING)
    return headers


def _ornament_values(i: int) -> list:
    """Distinct per ornament, so a mis-mapped column is unmistakable."""
    return [
        10.0 + i,        # Gross Wt.
        0.1 * i,         # Stone Wt.
        9.0 + i,         # Net Wt.
        22,              # Karat
        0.916,           # Purity % %
        8.0 + i,         # Net weight after Purity %
    ]


def _write_original_wide(path: str, ornaments: int, loans: int = 1) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_original_layout_headers(ornaments))
    for loan in range(1, loans + 1):
        row = [loan, "Hanamkonda", "WARANGAL", "05/03/2026", f"Customer {loan}",
               f"LN{loan:04d}", "secured", "Partner", f"CID{loan}"]
        for i in range(1, ornaments + 1):
            row.append(f"Ornament {i}")
            row.extend(_ornament_values(i))
        row.extend([100000 + loan, "Auditor", "Appraiser", "Vault", None,
                    f"PKT{loan}", f"PKT{loan}"])
        ws.append(row)
    wb.save(path)
    return path


def _convert(path: str):
    service = ArvogService()
    sheet, header_row = service.detect_raw_excel(path)
    assert sheet is not None, "fixture is not recognised as a wide master sheet"
    return service.convert_raw_to_tall(path, sheet, header_row, log_func=lambda m: None)


class TestOriginalLayoutIsReadCorrectly:
    def test_one_row_per_ornament(self, tmp_path):
        df = _convert(_write_original_wide(str(tmp_path / "w.xlsx"), ornaments=4))
        assert len(df) == 4
        assert df["Jewellery No."].tolist() == [f"Jewellery{i}" for i in range(1, 5)]

    def test_every_ornament_keeps_its_own_weights(self, tmp_path):
        """The mapping bug this suite guards against shows up here first."""
        df = _convert(_write_original_wide(str(tmp_path / "w.xlsx"), ornaments=6))
        assert df["Gross Wt."].tolist() == [10.0 + i for i in range(1, 7)]
        assert df["Net Wt."].tolist() == [9.0 + i for i in range(1, 7)]
        assert df["Net weight after Purity %"].tolist() == [8.0 + i for i in range(1, 7)]
        assert df["Stone Wt."].tolist() == pytest.approx([0.1 * i for i in range(1, 7)])

    def test_the_twentieth_ornament_is_read(self, tmp_path):
        """Its purity column is spelled 'Final net weight after purity %'."""
        df = _convert(_write_original_wide(str(tmp_path / "w.xlsx"), ornaments=20))
        assert len(df) == 20
        assert df["Gross Wt."].iloc[19] == 30.0
        assert df["Net weight after Purity %"].iloc[19] == 28.0

    def test_loan_fields_only_on_the_first_ornament_row(self, tmp_path):
        df = _convert(_write_original_wide(str(tmp_path / "w.xlsx"), ornaments=3))
        assert df["Customer Name"].iloc[0] == "Customer 1"
        # Continuation rows come back as NaN rather than None — pandas' doing,
        # pinned here because the rebuild has to recognise them as blanks.
        assert pd.isna(df["Customer Name"].iloc[1])
        assert pd.isna(df["Customer Name"].iloc[2])
        assert df["Loan Amount"].iloc[0] == 100001
        assert pd.isna(df["Loan Amount"].iloc[1])

    def test_several_loans(self, tmp_path):
        df = _convert(_write_original_wide(str(tmp_path / "w.xlsx"), ornaments=2, loans=3))
        assert len(df) == 6
        names = df["Customer Name"].tolist()
        assert [n if isinstance(n, str) else None for n in names] == [
            "Customer 1", None, "Customer 2", None, "Customer 3", None,
        ]

    def test_the_tall_column_order(self, tmp_path):
        df = _convert(_write_original_wide(str(tmp_path / "w.xlsx"), ornaments=2))
        assert list(df.columns) == LEADING + [
            "Jewellery No.", "Jewellery1", "Gross Wt.", "Stone Wt.", "Net Wt.",
            "Karat", "Purity % %", "Net weight after Purity %",
        ] + TRAILING

    def test_blank_and_dashed_ornaments_are_skipped(self, tmp_path):
        path = str(tmp_path / "w.xlsx")
        _write_original_wide(path, ornaments=3)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        ws.cell(row=2, column=headers.index("Jewellery2") + 1).value = "-"
        ws.cell(row=2, column=headers.index("Jewellery3") + 1).value = None
        wb.save(path)

        df = _convert(path)
        assert len(df) == 1
        assert df["Jewellery1"].tolist() == ["Ornament 1"]

    def test_absent_optional_columns_are_dropped_from_the_output(self, tmp_path):
        """Only the loan columns actually present in the file are emitted."""
        path = str(tmp_path / "w.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        headers = ["SR No.", "Branch", "Customer Name", "Loan No.",
                   "Jewellery1", "Gross Wt.", "Stone Wt.", "Net Wt.", "Karat",
                   "Purity % %", "Net weight after Purity %",
                   "Jewellery2", "Gross Wt.", "Stone Wt.", "Net Wt.", "Karat",
                   "Purity % %", "Net weight after Purity %",
                   "Packets Number"]
        ws.append(headers)
        ws.append([1, "Hanamkonda", "Ravi", "LN1",
                   "Bangle", 12.0, 0.0, 12.0, 22, 0.916, 11.0,
                   "Chain", 6.0, 0.5, 5.5, 22, 0.916, 5.0,
                   "PKT1"])
        wb.save(path)

        df = _convert(path)
        assert "Region" not in df.columns
        assert "Auditor Name" not in df.columns
        assert df["Gross Wt."].tolist() == [12.0, 6.0]
        assert df["Net Wt."].tolist() == [12.0, 5.5]
