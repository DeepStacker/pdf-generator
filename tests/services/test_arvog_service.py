"""Tests for the ArvogService class API."""

import pandas as pd
import pytest

from audit_engine.services.arvog import ArvogService


def test_service_can_be_instantiated():
    service = ArvogService()
    assert service is not None
    assert service.name == "Arvog Bank"


def test_static_normalize_columns():
    result = ArvogService.normalize_columns(["  Foo ", "BAR "])
    assert result == ["foo", "bar"]


def test_static_format_value():
    assert ArvogService.format_value(None) == ""
    assert ArvogService.format_value(pd.NA) == ""
    assert ArvogService.format_value(5.0) == "5"
    assert ArvogService.format_value(5.5) == "5.5"
    assert ArvogService.format_value("hello") == "hello"
    assert ArvogService.format_value("3.0") == "3"


def test_static_make_cell():
    from reportlab.lib.styles import ParagraphStyle
    style = ParagraphStyle("test", fontSize=10)
    cell = ArvogService.make_cell("hello", style)
    assert cell is not None
    assert "hello" in str(cell)
    empty = ArvogService.make_cell(None, style)
    assert empty == ""


def test_required_columns_present():
    assert len(ArvogService.REQUIRED_COLUMNS) == 11
    assert "SR No." in ArvogService.REQUIRED_COLUMNS
    assert "Branch" in ArvogService.REQUIRED_COLUMNS


def test_detect_raw_excel_nonexistent():
    service = ArvogService()
    sheet, row = service.detect_raw_excel("/nonexistent/file.xlsx")
    assert sheet is None
    assert row is None


def test_clean_dataframe_raises_on_missing(tmp_path):
    service = ArvogService()
    df = pd.DataFrame({"Foo": [1]})
    import pytest
    with pytest.raises(Exception, match="Missing columns"):
        service.clean_dataframe(df)


def test_service_process_excel_handles_missing_file(tmp_path):
    service = ArvogService()
    import pytest
    with pytest.raises(Exception):
        service.process_excel("/nonexistent/file.xlsx", str(tmp_path))


# =========================================================
# PDF table layout
# =========================================================

class TestPdfLayout:
    """The headers are drawn as plain strings, so reportlab does not wrap them:
    a column narrower than its own heading bleeds into its neighbour. These
    tests pin both the fit and the total, so retuning one column cannot quietly
    break another."""

    def _line_width(self, line):
        from reportlab.pdfbase.pdfmetrics import stringWidth

        from audit_engine.services.arvog import HEADER_FONT, HEADER_FONT_SIZE
        return stringWidth(line, HEADER_FONT, HEADER_FONT_SIZE)

    def test_one_width_per_header(self):
        from audit_engine.services.arvog import PDF_COLUMN_WIDTHS, PDF_HEADERS
        assert len(PDF_COLUMN_WIDTHS) == len(PDF_HEADERS)

    def test_every_header_fits_its_column(self):
        from audit_engine.services.arvog import (
            CELL_HORIZONTAL_PADDING,
            PDF_COLUMN_WIDTHS,
            PDF_HEADERS,
        )
        clipped = []
        for header, width in zip(PDF_HEADERS, PDF_COLUMN_WIDTHS):
            available = width - 2 * CELL_HORIZONTAL_PADDING
            needed = max(self._line_width(line) for line in header.split("\n"))
            if needed > available:
                clipped.append(f"{header!r} needs {needed:.2f}pt, has {available:.2f}pt")
        assert not clipped, "headers wider than their column: " + "; ".join(clipped)

    def test_headers_keep_a_little_breathing_room(self):
        """Fitting to the last fraction of a point is what forced the previous
        layout to abbreviate a heading."""
        from audit_engine.services.arvog import (
            CELL_HORIZONTAL_PADDING,
            PDF_COLUMN_WIDTHS,
            PDF_HEADERS,
        )
        for header, width in zip(PDF_HEADERS, PDF_COLUMN_WIDTHS):
            available = width - 2 * CELL_HORIZONTAL_PADDING
            needed = max(self._line_width(line) for line in header.split("\n"))
            assert available - needed >= 4.0, f"{header!r} is cramped"

    def test_columns_still_fill_the_printable_width(self):
        from audit_engine.services.arvog import PDF_COLUMN_WIDTHS, PRINTABLE_WIDTH
        total = sum(PDF_COLUMN_WIDTHS)
        assert total <= PRINTABLE_WIDTH
        assert total > PRINTABLE_WIDTH - 1.0

    def test_the_truncated_heading_is_spelled_out(self):
        from audit_engine.services.arvog import PDF_HEADERS
        assert "Count of\nOrnament" in PDF_HEADERS
        assert not any("Ornamen\n" in h or h.endswith("Ornamen") for h in PDF_HEADERS)

    def test_data_columns_line_up_with_the_filled_half(self):
        from audit_engine.services.arvog import (
            AUDIT_BLOCK_FIRST_COLUMN,
            PDF_DATA_COLUMNS,
            PDF_HEADERS,
        )
        assert len(PDF_DATA_COLUMNS) == AUDIT_BLOCK_FIRST_COLUMN
        assert len(PDF_HEADERS) > AUDIT_BLOCK_FIRST_COLUMN

    def test_every_pdf_data_column_has_a_declared_type(self):
        from audit_engine.services.arvog import PDF_DATA_COLUMNS
        from audit_engine.services.arvog_columns import kind_for
        assert all(kind_for(c) is not None for c in PDF_DATA_COLUMNS)


# =========================================================
# Wide -> tall conversion, end to end
# =========================================================

WIDE_HEADERS = [
    "SR No.", "Branch", "Region", "Date", "Customer Name", "Loan No.",
    "Jewellery1", "Gross Wt.", "Stone Wt.", "Net Wt.", "Karat", "Purity % %",
    "Jewellery2", "Gross Wt.", "Stone Wt.", "Net Wt.", "Karat", "Purity % %",
    "Loan Amount", "Auditor Name", "Appraiser Name", "Packets Location",
    "Cluster Manager", "Packets Number", "After audit packet number",
]


def _make_wide_xlsx(path, region="WARANGAL", cluster_manager=None):
    """A two-loan Arvog master sheet in the raw wide layout."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(WIDE_HEADERS)
    ws.append([
        1, "Hanamkonda", region, "05/03/2026", "  Ravi   Kumar ", 123456.0,
        "Bangle", 12.345, 0.0, 12.345, 22.0, 91.6,
        "Chain", 2, 0.5, 1.5, "22k", 91.6,
        125000, "Auditor A", "Appraiser B", "Strong Room",
        cluster_manager, 90210.0, 45.0,
    ])
    ws.append([
        2, "Hanamkonda", region, "06/03/2026", "Sita", 123457.0,
        "Ring", 3.0, 0.0, 3.0, 18.0, 75.0,
        "-", None, None, None, None, None,
        50000, "Auditor A", "Appraiser B", "Strong Room",
        cluster_manager, 90211.0, 46.0,
    ])
    wb.save(path)
    return path


class TestWideToTallConversion:
    def _convert(self, path, **kwargs):
        from audit_engine.services.arvog import ArvogService
        service = ArvogService()
        sheet, header_row = service.detect_raw_excel(path)
        assert sheet is not None, "fixture is not recognised as a raw wide sheet"
        return service.convert_raw_to_tall(path, sheet, header_row, log_func=lambda m: None, **kwargs)

    def test_each_ornament_becomes_a_row(self, tmp_path):
        df = self._convert(_make_wide_xlsx(str(tmp_path / "wide.xlsx")))
        # Loan 1 has two ornaments, loan 2 has one ("-" is skipped).
        assert len(df) == 3
        assert df["Jewellery No."].tolist() == ["Jewellery1", "Jewellery2", "Jewellery1"]
        assert df["Jewellery1"].tolist() == ["Bangle", "Chain", "Ring"]

    def test_only_the_first_ornament_row_carries_the_loan_fields(self, tmp_path):
        df = self._convert(_make_wide_xlsx(str(tmp_path / "wide.xlsx")))
        assert df["Customer Name"].iloc[1] is None or pd.isna(df["Customer Name"].iloc[1])

    def test_region_cluster_manager_comes_from_config(self, tmp_path):
        df = self._convert(_make_wide_xlsx(str(tmp_path / "wide.xlsx"), region="WARANGAL"))
        from audit_engine.utils.config import arvog as arvog_config
        assert df["Cluster Manager"].iloc[0] == arvog_config.cluster_manager_for("WARANGAL")

    def test_unconfigured_region_gets_no_manager(self, tmp_path):
        df = self._convert(_make_wide_xlsx(str(tmp_path / "wide.xlsx"), region="NOWHERE"))
        assert df["Cluster Manager"].iloc[0] is None

    def test_explicit_cluster_manager_wins(self, tmp_path):
        df = self._convert(
            _make_wide_xlsx(str(tmp_path / "wide.xlsx"), region="WARANGAL"),
            cluster_manager="Explicit Name",
        )
        assert df["Cluster Manager"].iloc[0] == "Explicit Name"

    def test_a_manager_already_in_the_sheet_is_kept(self, tmp_path):
        df = self._convert(_make_wide_xlsx(str(tmp_path / "wide.xlsx"), cluster_manager="From Sheet"))
        assert df["Cluster Manager"].iloc[0] == "From Sheet"


class TestSchemaReference:
    """The reference sheet used to be probed implicitly at
    ``manual_converted/<filename>`` relative to the working directory, so the
    same input converted differently depending on where you ran it from."""

    def _wide(self, tmp_path):
        return _make_wide_xlsx(str(tmp_path / "wide.xlsx"))

    def test_a_manual_converted_dir_in_the_cwd_is_ignored(self, tmp_path, monkeypatch):
        import openpyxl

        from audit_engine.services.arvog import ArvogService

        wide = self._wide(tmp_path)
        stray = tmp_path / "manual_converted"
        stray.mkdir()
        wb = openpyxl.Workbook()
        wb.active.append(["Totally", "Different", "Schema"])
        wb.save(str(stray / "wide.xlsx"))

        monkeypatch.chdir(tmp_path)
        service = ArvogService()
        sheet, header_row = service.detect_raw_excel(wide)
        df = service.convert_raw_to_tall(wide, sheet, header_row, log_func=lambda m: None)

        assert "Totally" not in df.columns
        assert "Customer Name" in df.columns

    def test_an_explicit_reference_sets_the_columns(self, tmp_path):
        import openpyxl

        from audit_engine.services.arvog import ArvogService

        wide = self._wide(tmp_path)
        ref = tmp_path / "reference.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["SR No.", "Customer Name", "Jewellery No.", "Jewellery1", "Gross Wt."])
        wb.save(str(ref))

        service = ArvogService()
        sheet, header_row = service.detect_raw_excel(wide)
        df = service.convert_raw_to_tall(
            wide, sheet, header_row, log_func=lambda m: None, schema_reference=str(ref),
        )
        assert list(df.columns) == ["SR No.", "Customer Name", "Jewellery No.", "Jewellery1", "Gross Wt."]

    def test_a_missing_reference_is_reported_and_does_not_raise(self, tmp_path):
        from audit_engine.services.arvog import ArvogService
        messages = []
        service = ArvogService()
        wide = self._wide(tmp_path)
        sheet, header_row = service.detect_raw_excel(wide)
        df = service.convert_raw_to_tall(
            wide, sheet, header_row, log_func=messages.append,
            schema_reference=str(tmp_path / "absent.xlsx"),
        )
        assert "Customer Name" in df.columns
        assert any("not found" in m for m in messages)


def _converted_sheet(tmp_path, output_format="EXCEL ONLY"):
    """Run a conversion and hand back the saved worksheet."""
    import openpyxl

    from audit_engine.services.arvog import ArvogService
    out_dir = tmp_path / "out"
    ArvogService().process_excel(
        input_excel=_make_wide_xlsx(str(tmp_path / "wide.xlsx")),
        output_dir=str(out_dir),
        log_func=lambda m: None,
        output_format=output_format,
    )
    return openpyxl.load_workbook(str(out_dir / "wide.xlsx")).active, out_dir


def _headers(ws):
    from audit_engine.services.arvog import EXCEL_HEADER_ROW
    return [c.value for c in ws[EXCEL_HEADER_ROW]]


def _first_data_row(ws):
    from audit_engine.services.arvog import EXCEL_HEADER_ROW
    return EXCEL_HEADER_ROW + 1


def _cell(ws, header, row_offset=0):
    from audit_engine.services.arvog import EXCEL_HEADER_ROW
    col = _headers(ws).index(header) + 1
    return ws.cell(row=EXCEL_HEADER_ROW + 1 + row_offset, column=col)


class TestAuditBlock:
    """The auditor's hand-filled tally closes every converted sheet. It used to
    arrive only from a reference workbook found relative to the working
    directory, so whether the report had one at all depended on where the
    process was started."""

    def test_the_block_is_present_and_last(self, tmp_path):
        from audit_engine.services.arvog import EXCEL_AUDIT_BLOCK_COLUMNS
        ws, _ = _converted_sheet(tmp_path)
        assert _headers(ws)[-len(EXCEL_AUDIT_BLOCK_COLUMNS):] == EXCEL_AUDIT_BLOCK_COLUMNS

    def test_the_required_headings_are_all_there_in_order(self, tmp_path):
        ws, _ = _converted_sheet(tmp_path)
        headers = [str(h) for h in _headers(ws)]
        required = [
            "Jewellery No.", "Jewellery1", "Gross Wt.", "Stone Wt.", "Net Wt.", "Karat",
            "Purity % %", "Net weight after Purity %", "Loan Amount", "Auditor Name",
            "Appraiser Name", "Packets Location", "Cluster Manager", "Packets Number",
            "After audit packet number", "Count of Ornament", "Gross Wt.", "Stone Wt",
            "Karat", "Purity%", "Net Weight", "Remarks",
        ]
        start = headers.index("Jewellery No.")
        assert headers[start:start + len(required)] == required

    def test_repeated_headings_are_not_collapsed(self, tmp_path):
        """"Gross Wt." and "Karat" appear on both sides of the sheet. Building
        rows through a dict silently merged them."""
        ws, _ = _converted_sheet(tmp_path)
        headers = [str(h) for h in _headers(ws)]
        assert headers.count("Gross Wt.") == 2
        assert headers.count("Karat") == 2

    def test_the_block_is_left_blank_for_the_auditor(self, tmp_path):
        from audit_engine.services.arvog import EXCEL_AUDIT_BLOCK_COLUMNS
        ws, _ = _converted_sheet(tmp_path)
        span_start = len(_headers(ws)) - len(EXCEL_AUDIT_BLOCK_COLUMNS) + 1
        for row in range(_first_data_row(ws), ws.max_row + 1):
            for col in range(span_start, ws.max_column + 1):
                assert ws.cell(row=row, column=col).value is None

    def test_sumeru_banners_the_block(self, tmp_path):
        from audit_engine.services.arvog import (
            EXCEL_AUDIT_BANNER,
            EXCEL_AUDIT_BLOCK_COLUMNS,
            EXCEL_BANNER_ROW,
        )
        ws, _ = _converted_sheet(tmp_path)
        span_start = len(_headers(ws)) - len(EXCEL_AUDIT_BLOCK_COLUMNS) + 1
        assert ws.cell(row=EXCEL_BANNER_ROW, column=span_start).value == EXCEL_AUDIT_BANNER

    def test_the_banner_is_merged_across_the_whole_block(self, tmp_path):
        from audit_engine.services.arvog import EXCEL_AUDIT_BLOCK_COLUMNS, EXCEL_BANNER_ROW
        ws, _ = _converted_sheet(tmp_path)
        span_start = len(_headers(ws)) - len(EXCEL_AUDIT_BLOCK_COLUMNS) + 1
        merged = [str(r) for r in ws.merged_cells.ranges]
        assert any(
            r.min_row == EXCEL_BANNER_ROW and r.min_col == span_start and r.max_col == ws.max_column
            for r in ws.merged_cells.ranges
        ), f"banner not merged across the block: {merged}"

    def test_nothing_is_banner_ed_over_the_data_half(self, tmp_path):
        from audit_engine.services.arvog import EXCEL_AUDIT_BLOCK_COLUMNS, EXCEL_BANNER_ROW
        ws, _ = _converted_sheet(tmp_path)
        span_start = len(_headers(ws)) - len(EXCEL_AUDIT_BLOCK_COLUMNS) + 1
        for col in range(1, span_start):
            assert ws.cell(row=EXCEL_BANNER_ROW, column=col).value is None

    def test_the_block_is_yellow_and_the_data_headings_are_not(self, tmp_path):
        from audit_engine.services.arvog import (
            _YELLOW,
            EXCEL_AUDIT_BLOCK_COLUMNS,
            EXCEL_HEADER_ROW,
        )
        ws, _ = _converted_sheet(tmp_path)
        span_start = len(_headers(ws)) - len(EXCEL_AUDIT_BLOCK_COLUMNS) + 1
        assert ws.cell(row=EXCEL_HEADER_ROW, column=span_start).fill.fgColor.rgb == _YELLOW
        assert ws.cell(row=EXCEL_HEADER_ROW, column=1).fill.fgColor.rgb != _YELLOW

    def test_the_block_columns_carry_their_own_number_formats(self, tmp_path):
        ws, _ = _converted_sheet(tmp_path)
        assert _cell(ws, "Count of Ornament").number_format == "0"
        assert _cell(ws, "Net Weight").number_format == "0.00"
        assert _cell(ws, "Purity%").number_format == "0.00%"
        assert _cell(ws, "Remarks").number_format == "@"

    def test_appending_twice_does_not_double_the_block(self, tmp_path):
        import pandas as pd

        from audit_engine.services.arvog import EXCEL_AUDIT_BLOCK_COLUMNS, ArvogService
        service = ArvogService()
        once = service.append_audit_block(pd.DataFrame({"SR No.": [1]}))
        twice = service.append_audit_block(once)
        assert list(twice.columns) == list(once.columns)
        assert list(once.columns)[-len(EXCEL_AUDIT_BLOCK_COLUMNS):] == EXCEL_AUDIT_BLOCK_COLUMNS

    def test_an_empty_frame_still_gets_the_block(self, tmp_path):
        import pandas as pd

        from audit_engine.services.arvog import EXCEL_AUDIT_BLOCK_COLUMNS, ArvogService
        out = ArvogService().append_audit_block(pd.DataFrame({"SR No.": []}))
        assert list(out.columns)[-len(EXCEL_AUDIT_BLOCK_COLUMNS):] == EXCEL_AUDIT_BLOCK_COLUMNS
        assert len(out) == 0

    def test_the_pdf_frame_is_untouched_by_the_block(self, tmp_path):
        """Repeated headings would make row.get("Gross Wt.") return a Series."""
        ws, out_dir = _converted_sheet(tmp_path, output_format="BOTH")
        assert (out_dir / "Hanamkonda.pdf").exists()
        assert (out_dir / "Hanamkonda.pdf").stat().st_size > 0


class TestProcessExcelOutputs:
    def test_produces_a_pdf_per_branch(self, tmp_path):
        _, out_dir = _converted_sheet(tmp_path, output_format="BOTH")
        assert (out_dir / "Hanamkonda.pdf").exists()

    def test_excel_number_formats_follow_the_column_types(self, tmp_path):
        ws, _ = _converted_sheet(tmp_path)
        assert _cell(ws, "Customer Name").number_format == "@"
        assert _cell(ws, "Loan No.").number_format == "@"
        assert _cell(ws, "Packets Number").number_format == "@"
        assert _cell(ws, "Loan Amount").number_format == "#,##0.00"
        assert _cell(ws, "Purity % %").number_format == "0.00%"
        assert _cell(ws, "Date").number_format == "dd-mmm-yy"

    def test_a_packet_number_is_not_formatted_as_a_date(self, tmp_path):
        ws, _ = _converted_sheet(tmp_path)
        assert _cell(ws, "After audit packet number").number_format == "@"

    def test_values_are_normalised_in_the_saved_excel(self, tmp_path):
        ws, _ = _converted_sheet(tmp_path)
        assert _cell(ws, "Customer Name").value == "Ravi Kumar"
        assert _cell(ws, "Loan No.").value == "123456"
        assert _cell(ws, "Gross Wt.").value == 12.345
        assert _cell(ws, "Purity % %").value == pytest.approx(0.916)

    def test_karat_column_with_mixed_conventions_is_stored_as_text(self, tmp_path):
        ws, _ = _converted_sheet(tmp_path)
        assert _cell(ws, "Karat").number_format == "@"

    def test_excel_only_skips_the_pdfs(self, tmp_path):
        _, out_dir = _converted_sheet(tmp_path, output_format="EXCEL ONLY")
        assert not list(out_dir.glob("*.pdf"))
        assert (out_dir / "wide.xlsx").exists()

    def test_pdf_only_skips_the_excel(self, tmp_path):
        from audit_engine.services.arvog import ArvogService
        out_dir = tmp_path / "out"
        ArvogService().process_excel(
            input_excel=_make_wide_xlsx(str(tmp_path / "wide.xlsx")),
            output_dir=str(out_dir), log_func=lambda m: None, output_format="PDF ONLY",
        )
        assert (out_dir / "Hanamkonda.pdf").exists()
        assert not (out_dir / "wide.xlsx").exists()


class TestExcelPresentation:
    """It used to come out with no alignment, no borders and every column at
    Excel's 8.43 default, which cut the headings."""

    def test_every_cell_is_centred(self, tmp_path):
        from audit_engine.services.arvog import EXCEL_HEADER_ROW
        ws, _ = _converted_sheet(tmp_path)
        for row in ws.iter_rows(min_row=EXCEL_HEADER_ROW, max_row=ws.max_row):
            for cell in row:
                assert cell.alignment.horizontal == "center", f"{cell.coordinate} not centred"
                assert cell.alignment.vertical == "center", f"{cell.coordinate} not centred"

    def test_every_cell_is_bordered_on_all_four_sides(self, tmp_path):
        from audit_engine.services.arvog import EXCEL_HEADER_ROW
        ws, _ = _converted_sheet(tmp_path)
        for row in ws.iter_rows(min_row=EXCEL_HEADER_ROW, max_row=ws.max_row):
            for cell in row:
                sides = (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom)
                assert all(s.style for s in sides), f"{cell.coordinate} is not boxed"

    def test_blank_cells_inside_the_table_are_boxed_too(self, tmp_path):
        ws, _ = _converted_sheet(tmp_path)
        blanks = [c for row in ws.iter_rows(min_row=_first_data_row(ws)) for c in row if c.value is None]
        assert blanks, "fixture no longer has an empty cell to check"
        for cell in blanks:
            assert cell.border.left.style and cell.border.bottom.style

    def test_every_row_is_the_same_height(self, tmp_path):
        """No row carries an explicit height, so banner, heading and data all
        come out at the sheet default."""
        ws, _ = _converted_sheet(tmp_path)
        heights = {
            row: ws.row_dimensions[row].height
            for row in range(1, ws.max_row + 1)
            if ws.row_dimensions[row].height is not None
        }
        assert not heights, f"rows given their own height: {heights}"

    def test_nothing_wraps(self, tmp_path):
        """A wrapping heading needs a taller row than the data beneath it."""
        ws, _ = _converted_sheet(tmp_path)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                assert not cell.alignment.wrap_text, f"{cell.coordinate} wraps"

    def test_each_column_fits_its_whole_heading_on_one_line(self, tmp_path):
        from openpyxl.utils import get_column_letter

        from audit_engine.services.arvog import EXCEL_MAX_COLUMN_WIDTH
        ws, _ = _converted_sheet(tmp_path)
        for col, header in enumerate(_headers(ws), start=1):
            width = ws.column_dimensions[get_column_letter(col)].width
            needed = len(str(header or ""))
            assert width >= min(needed, EXCEL_MAX_COLUMN_WIDTH), f"{header!r} is cut at {width}"

    def test_no_column_is_left_at_the_default_width(self, tmp_path):
        from openpyxl.utils import get_column_letter

        from audit_engine.services.arvog import EXCEL_MAX_COLUMN_WIDTH, EXCEL_MIN_COLUMN_WIDTH
        ws, _ = _converted_sheet(tmp_path)
        for c in range(1, ws.max_column + 1):
            width = ws.column_dimensions[get_column_letter(c)].width
            assert width is not None
            assert EXCEL_MIN_COLUMN_WIDTH <= width <= EXCEL_MAX_COLUMN_WIDTH

    def test_width_follows_the_displayed_text_not_the_stored_value(self, tmp_path):
        """A date is stored as 2026-03-05 00:00:00 but shown as 05-Mar-26."""
        from openpyxl.utils import get_column_letter
        ws, _ = _converted_sheet(tmp_path)
        col = get_column_letter([str(h) for h in _headers(ws)].index("Date") + 1)
        assert ws.column_dimensions[col].width <= len("05-Mar-26") + 3

    def test_a_formatted_amount_still_fits(self, tmp_path):
        """125000 is stored short but displays as 125,000.00."""
        from openpyxl.utils import get_column_letter
        ws, _ = _converted_sheet(tmp_path)
        col = get_column_letter([str(h) for h in _headers(ws)].index("Loan Amount") + 1)
        assert ws.column_dimensions[col].width >= len("125,000.00")
